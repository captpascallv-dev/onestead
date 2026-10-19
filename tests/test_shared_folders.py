"""共享文件夹：只给点名的伙伴只读搜索；不跟随链接、不碰登录资料和凭据文件、不让路径跳出文件夹。"""

import json
import os
import sqlite3
import sys
import threading
import time
from pathlib import Path
from types import SimpleNamespace

import pytest

from onestead import config, folders, research
from onestead.store import BaseError


@pytest.fixture(autouse=True)
def _stop_folder_index():
    folders.shutdown()
    yield
    folders.shutdown()


@pytest.fixture
def shared(tmp_path, monkeypatch):
    root = tmp_path / "调研"
    outside = tmp_path / "outside"
    (root / "报告").mkdir(parents=True)
    outside.mkdir()
    (outside / "机密.md").write_text("外面的暗号橙子", encoding="utf-8")
    (root / "报告" / "结论.md").write_text("市场结论：暗号蓝莓。", encoding="utf-8")
    (root / "报告" / "网页.html").write_text("<html><script>x=1</script><p>暗号&amp;葡萄</p></html>", encoding="utf-8")
    (root / "gbk.txt").write_bytes("旧编码暗号柠檬".encode("gb18030"))
    (root / "工具.py").write_text("暗号蓝莓", encoding="utf-8")
    (root / "secret-notes.md").write_text("暗号蓝莓", encoding="utf-8")
    (root / "app.local.json").write_text('{"x": "暗号蓝莓"}', encoding="utf-8")
    (root / ".env").write_text("暗号蓝莓", encoding="utf-8")
    (root / "带凭据.md").write_text("api_key = sk-abcdefghijklmnopqrstuvwxyz123456 暗号凭据", encoding="utf-8")
    browser = root / "采集" / "native" / "user-data"
    browser.mkdir(parents=True)
    (browser / "Local State").write_text("{}", encoding="utf-8")
    (browser / "notes.md").write_text("暗号蓝莓", encoding="utf-8")
    (root / "采集" / "chrome-profile").mkdir()
    (root / "采集" / "chrome-profile" / "a.md").write_text("暗号蓝莓", encoding="utf-8")
    (root / "排除我").mkdir()
    (root / "排除我" / "a.md").write_text("暗号蓝莓", encoding="utf-8")
    (root / "node_modules").mkdir()
    (root / "node_modules" / "a.md").write_text("暗号蓝莓", encoding="utf-8")
    linked = False
    if sys.platform == "win32":
        import _winapi

        _winapi.CreateJunction(str(outside), str(root / "跳出去"))
        linked = True
    else:
        os.symlink(outside, root / "跳出去")
        linked = True
    rows = [{"key": "research", "name": "调研", "path": str(root), "readers": ["chatgpt"], "exclude": ["排除我"]}]
    monkeypatch.setenv("ONESTEAD_SHARED_FOLDERS", json.dumps(rows, ensure_ascii=False))
    return SimpleNamespace(root=root, outside=outside, linked=linked)


def _rels():
    folder = folders.folders_for("chatgpt")[0]
    return sorted(rel for rel, _, _ in folders.walk(folder))


def test_walk_keeps_only_whitelisted_files(shared):
    assert _rels() == ["gbk.txt", "带凭据.md", "报告/结论.md", "报告/网页.html"]


def test_only_readers_see_the_folder(shared):
    assert folders.search("chatgpt", ["蓝莓"])
    assert folders.search("claude", ["蓝莓"]) == []
    assert folders.search(None, ["蓝莓"]) == []
    assert folders.fetch("claude", "folder:research/报告/结论.md") is None


def test_search_hits_body_and_path(shared):
    titles = [row["title"] for row in folders.search("chatgpt", ["蓝莓"])]
    assert titles == ["调研/报告/结论.md"]
    assert [row["title"] for row in folders.search("chatgpt", ["报告", "葡萄"])] == ["调研/报告/网页.html"]
    assert [row["title"] for row in folders.search("chatgpt", ["柠檬"])] == ["调研/gbk.txt"]
    assert folders.search("chatgpt", ["橙子"]) == []


@pytest.mark.parametrize("rel", [
    "../outside/机密.md", "报告/../../outside/机密.md", "跳出去/机密.md", "/etc/passwd", "C:/Windows/win.ini",
    "报告\\结论.md", "工具.py", "secret-notes.md", "app.local.json", ".env", "排除我/a.md", "node_modules/a.md",
    "采集/native/user-data/notes.md", "采集/chrome-profile/a.md", "报告/./结论.md", "报告//结论.md",
])
def test_fetch_refuses_paths_outside_the_rules(shared, rel):
    assert folders.fetch("chatgpt", f"folder:research/{rel}") is None


def test_fetch_reads_and_strips_html(shared):
    row = folders.fetch("chatgpt", "folder:research/报告/网页.html")
    assert "暗号&葡萄" in row["text"] and "x=1" not in row["text"]


def test_office_files_are_extracted(shared):
    import docx
    import openpyxl
    from pptx import Presentation

    d = docx.Document()
    d.add_paragraph("Word 暗号西瓜")
    d.save(shared.root / "报告" / "a.docx")
    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[1])
    slide.shapes.title.text = "PPT 暗号桃子"
    p.save(shared.root / "报告" / "b.pptx")
    wb = openpyxl.Workbook()
    wb.active.append(["Excel", "暗号芒果"])
    wb.save(shared.root / "报告" / "c.xlsx")
    for word, name in (("西瓜", "a.docx"), ("桃子", "b.pptx"), ("芒果", "c.xlsx")):
        assert [r["title"] for r in folders.search("chatgpt", [word])] == [f"调研/报告/{name}"]
        assert word in folders.fetch("chatgpt", f"folder:research/报告/{name}")["text"]


def test_research_fetch_redacts_and_hides(shared):
    base = SimpleNamespace(_ctx=lambda: (SimpleNamespace(key="chatgpt", known=True, valid=True), [], []))
    out = research.fetch(base, "folder:research/带凭据.md", "https://base.example.test")
    assert "暗号凭据" in out.text and "sk-abcdefghijklmnopqrstuvwxyz123456" not in out.text
    assert out.url.startswith("https://base.example.test/#/folder/")
    with pytest.raises(BaseError):
        research.fetch(base, "folder:research/../outside/机密.md", "https://base.example.test")
    other = SimpleNamespace(_ctx=lambda: (SimpleNamespace(key="claude", known=True, valid=True), [], []))
    with pytest.raises(BaseError):
        research.fetch(other, "folder:research/报告/结论.md", "https://base.example.test")


@pytest.mark.parametrize("rows", [
    {"key": "x"}, [{"key": "x", "name": "n", "path": "p"}],
    [{"key": "x", "name": "n", "path": "p", "readers": []}],
    [{"key": "X Y", "name": "n", "path": "p", "readers": ["a"]}],
    [{"key": "x", "name": "n", "path": "p", "readers": ["a"]}, {"key": "x", "name": "m", "path": "q", "readers": ["a"]}],
    [{"key": "x", "name": "n", "path": "p", "readers": ["a"], "exclude": "nope"}],
])
def test_bad_config_is_rejected(monkeypatch, rows):
    monkeypatch.setenv("ONESTEAD_SHARED_FOLDERS", json.dumps(rows))
    with pytest.raises(BaseError):
        config.shared_folders()


def _count_extracts(monkeypatch):
    calls = []
    real = folders.text_of

    def wrap(path, st=None):
        calls.append(path.name)
        return real(path, st)

    monkeypatch.setattr(folders, "text_of", wrap)
    return calls


def _index_path():
    return config.state_dir() / folders.INDEX_FILENAME


def _remove_index():
    folders.close()
    path = _index_path()
    for suffix in ("-wal", "-shm", ""):
        candidate = Path(str(path) + suffix)
        if candidate.exists():
            candidate.unlink()


def _rows(folder_key=None):
    conn = sqlite3.connect(_index_path())
    try:
        if folder_key is None:
            return conn.execute("SELECT folder_key, rel FROM files").fetchall()
        return conn.execute("SELECT rel FROM files WHERE folder_key=?", (folder_key,)).fetchall()
    finally:
        conn.close()


def test_index_lives_in_the_state_dir(shared):
    assert not hasattr(folders, "_cache")
    folders.sync()
    path = _index_path()
    assert path.is_file()
    assert path.parent == Path(os.environ["COMPANY_BASE_STATE_DIR"])
    assert path.name == "shared-folders.sqlite"


def test_reopen_does_not_extract_unchanged_files(shared, monkeypatch):
    calls = _count_extracts(monkeypatch)
    folders.sync()
    first = len(calls)
    assert first >= 1
    folders.close()
    assert "蓝莓" in folders.fetch("chatgpt", "folder:research/报告/结论.md")["text"]
    folders.sync()
    assert len(calls) == first


def test_sync_reextracts_only_the_changed_file(shared, monkeypatch):
    calls = _count_extracts(monkeypatch)
    folders.sync()
    before = len(calls)
    target = shared.root / "报告" / "结论.md"
    target.write_text("市场结论：暗号蓝莓改了。", encoding="utf-8")
    folders.sync()
    assert calls[before:] == ["结论.md"]
    assert "改了" in folders.fetch("chatgpt", "folder:research/报告/结论.md")["text"]
    assert len(calls) == before + 1


def test_deleted_file_drops_out_of_search(shared):
    folders.sync()
    (shared.root / "报告" / "结论.md").unlink()
    folders.sync()
    assert folders.search("chatgpt", ["蓝莓"]) == []
    assert ("research", "报告/结论.md") not in _rows()


def test_new_exclude_drops_indexed_rows(shared, monkeypatch):
    folders.sync()
    assert ("research", "报告/网页.html") in _rows()
    rows = json.loads(os.environ["ONESTEAD_SHARED_FOLDERS"])
    rows[0]["exclude"] = ["排除我", "报告"]
    monkeypatch.setenv("ONESTEAD_SHARED_FOLDERS", json.dumps(rows, ensure_ascii=False))
    folders.sync()
    assert folders.search("chatgpt", ["葡萄"]) == []
    assert all(not rel.startswith("报告/") for _, rel in _rows())


def test_removed_folder_is_dropped(shared, tmp_path, monkeypatch):
    other = tmp_path / "other"
    other.mkdir()
    (other / "a.txt").write_text("暗号草莓", encoding="utf-8")
    rows = json.loads(os.environ["ONESTEAD_SHARED_FOLDERS"])
    rows.append({"key": "other", "name": "另一", "path": str(other), "readers": ["chatgpt"], "exclude": []})
    monkeypatch.setenv("ONESTEAD_SHARED_FOLDERS", json.dumps(rows, ensure_ascii=False))
    folders.sync()
    assert [row["title"] for row in folders.search("chatgpt", ["草莓"])] == ["另一/a.txt"]
    rows.pop()
    monkeypatch.setenv("ONESTEAD_SHARED_FOLDERS", json.dumps(rows, ensure_ascii=False))
    folders.sync()
    assert folders.search("chatgpt", ["草莓"]) == []
    assert {key for key, _ in _rows()} == {"research"}


def test_short_chinese_and_longer_terms(shared):
    (shared.root / "明天.txt").write_text("明天需要补货。", encoding="utf-8")
    (shared.root / "单据.txt").write_text("请看补货单详情。", encoding="utf-8")
    (shared.root / "登记.txt").write_text("这份是补货登记。", encoding="utf-8")
    assert {row["title"] for row in folders.search("chatgpt", ["补货"])} == {
        "调研/明天.txt", "调研/单据.txt", "调研/登记.txt",
    }
    assert [row["title"] for row in folders.search("chatgpt", ["补货单"])] == ["调研/单据.txt"]


def test_path_hits_rank_before_body_and_case_is_ignored(shared):
    (shared.root / "kalamata.txt").write_text("无", encoding="utf-8")
    (shared.root / "notes.txt").write_text("see kalamata here", encoding="utf-8")
    (shared.root / "蓝莓.txt").write_text("没有别的", encoding="utf-8")
    (shared.root / "正文.txt").write_text("这里有蓝莓", encoding="utf-8")
    assert [row["title"] for row in folders.search("chatgpt", ["Kalamata"])] == [
        "调研/kalamata.txt", "调研/notes.txt",
    ]
    assert [row["title"] for row in folders.search("chatgpt", ["蓝莓"])] == [
        "调研/蓝莓.txt", "调研/报告/结论.md", "调研/正文.txt",
    ]


def test_like_treats_percent_and_underscore_as_literal(tmp_path, monkeypatch):
    root = tmp_path / "box"
    root.mkdir()
    (root / "a.txt").write_text("keep a_b please", encoding="utf-8")
    (root / "b.txt").write_text("keep axb please", encoding="utf-8")
    (root / "c.txt").write_text("100% done", encoding="utf-8")
    monkeypatch.setenv("ONESTEAD_SHARED_FOLDERS", json.dumps([{
        "key": "box", "name": "盒子", "path": str(root), "readers": ["chatgpt"], "exclude": [],
    }], ensure_ascii=False))
    assert [row["title"] for row in folders.search("chatgpt", ["a_b"])] == ["盒子/a.txt"]
    assert [row["title"] for row in folders.search("chatgpt", ["%"])] == ["盒子/c.txt"]


def test_search_during_sync_does_not_error(shared, monkeypatch):
    started = threading.Event()
    release = threading.Event()
    real = folders.text_of

    def slow(path, st=None):
        started.set()
        assert release.wait(5)
        return real(path, st)

    monkeypatch.setattr(folders, "text_of", slow)

    def run():
        try:
            folders.sync()
        finally:
            folders.close()

    thread = threading.Thread(target=run)
    thread.start()
    assert started.wait(5)
    try:
        rows = folders.search("chatgpt", ["蓝莓"])
        assert isinstance(rows, list)
    finally:
        release.set()
    thread.join(10)
    assert not thread.is_alive()
    assert [row["title"] for row in folders.search("chatgpt", ["蓝莓"])] == ["调研/报告/结论.md"]


def test_only_one_sync_runs_at_a_time(shared, monkeypatch):
    current = {"n": 0, "max": 0}
    guard = threading.Lock()
    real = folders.text_of

    def wrap(path, st=None):
        with guard:
            current["n"] += 1
            current["max"] = max(current["max"], current["n"])
        try:
            time.sleep(0.05)
            return real(path, st)
        finally:
            with guard:
                current["n"] -= 1

    monkeypatch.setattr(folders, "text_of", wrap)

    def run():
        try:
            folders.sync()
        finally:
            folders.close()

    threads = [threading.Thread(target=run) for _ in range(2)]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join(10)
    assert current["max"] == 1
    assert all(not thread.is_alive() for thread in threads)


def test_deleted_index_is_rebuilt(shared, monkeypatch):
    calls = _count_extracts(monkeypatch)
    assert folders.search("chatgpt", ["蓝莓"])
    first = len(calls)
    _remove_index()
    assert [row["title"] for row in folders.search("chatgpt", ["蓝莓"])] == ["调研/报告/结论.md"]
    assert len(calls) == first * 2


def test_corrupt_index_is_rebuilt(shared):
    folders.sync()
    _remove_index()
    _index_path().write_bytes(b"this is not a database")
    assert [row["title"] for row in folders.search("chatgpt", ["蓝莓"])] == ["调研/报告/结论.md"]


def test_warm_syncs_again_after_the_interval(shared, monkeypatch):
    monkeypatch.setattr(folders, "SYNC_INTERVAL", 0.05)
    calls = {"n": 0}

    def fake(viewer_keys=None):
        calls["n"] += 1

    monkeypatch.setattr(folders, "sync", fake)
    started = time.monotonic()
    folders.warm()
    assert time.monotonic() - started < 1
    deadline = time.monotonic() + 3
    while calls["n"] < 2 and time.monotonic() < deadline:
        time.sleep(0.02)
    assert calls["n"] >= 2


def test_search_of_two_thousand_files_is_under_a_second(tmp_path, monkeypatch):
    root = tmp_path / "bulk"
    root.mkdir()
    for i in range(2000):
        (root / f"n{i:04d}.txt").write_text(f"记录 {i} 普通内容", encoding="utf-8")
    (root / "n0001.txt").write_text("这里有补货需求", encoding="utf-8")
    (root / "n0002.txt").write_text("唯一标记 kalamata 橄榄", encoding="utf-8")
    monkeypatch.setenv("ONESTEAD_SHARED_FOLDERS", json.dumps([{
        "key": "bulk", "name": "批量", "path": str(root), "readers": ["chatgpt"], "exclude": [],
    }], ensure_ascii=False))
    folders.sync()
    started = time.perf_counter()
    long_hits = folders.search("chatgpt", ["kalamata"])
    long_s = time.perf_counter() - started
    started = time.perf_counter()
    short_hits = folders.search("chatgpt", ["补货"])
    short_s = time.perf_counter() - started
    timing = Path(os.environ.get("TEMP", ".")) / "folder-index-search-seconds.txt"
    timing.write_text(f"long={long_s:.4f}\nshort={short_s:.4f}\n", encoding="utf-8")
    assert [row["title"] for row in long_hits] == ["批量/n0002.txt"]
    assert [row["title"] for row in short_hits] == ["批量/n0001.txt"]
    assert long_s < 1, long_s
    assert short_s < 1, short_s


def test_index_mode_is_private_on_posix(tmp_path, monkeypatch):
    path = tmp_path / "shared-folders.sqlite"
    path.write_bytes(b"")
    (tmp_path / "shared-folders.sqlite-wal").write_bytes(b"")
    modes = []
    monkeypatch.setattr(os, "name", "posix")

    def record(target, mode):
        modes.append((Path(target).name, mode))

    monkeypatch.setattr(os, "chmod", record)
    folders._restrict(path)
    assert any(str(target).endswith("shared-folders.sqlite") and mode == 0o600 for target, mode in modes)
    assert any(str(target).endswith("shared-folders.sqlite-wal") and mode == 0o600 for target, mode in modes)
