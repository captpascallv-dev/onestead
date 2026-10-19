"""共享文件夹：把主人在配置里点名的本机文件夹，只读地开放给指定伙伴搜索和读正文。

- 只认配置 shared_folders 里的文件夹，只给 readers 里的身份看；不进网页、同事视图和公开页面。
- 只读白名单类型的文件；跳过符号链接和目录联接、浏览器登录资料、带凭据字样的文件和目录。
- 读文件时从文件夹根逐级核对，路径里不许有 ..、盘符或链接，解析后的真实路径也必须还在文件夹里。
- PDF、Word、PPT、Excel 抽出文字；正文出去前仍走 safe_text 遮掉凭据和本机路径。
- 抽出的文字写在状态目录的 SQLite 索引里。原文件仍是唯一真相；索引删了或坏了按原文件重建。
"""

from __future__ import annotations

import html
import logging
import os
import re
import sqlite3
import threading
from fnmatch import fnmatch
from pathlib import Path

from . import config
from .i18n import text as tr

log = logging.getLogger(__name__)
logging.getLogger("pypdf").setLevel(logging.ERROR)  # 字体解析的提示每页一条，会刷满日志

TEXT_TYPES = {".md", ".txt", ".csv", ".json", ".html", ".htm"}
DOC_TYPES = {".pdf", ".docx", ".pptx", ".xlsx"}
MAX_TEXT_BYTES = 5 * 1024 * 1024
MAX_DOC_BYTES = 50 * 1024 * 1024
MAX_CHARS = 200_000  # 一份正文最多给这么多字，超出的截断并注明
SKIP_DIRS = {".git", ".svn", ".hg", "node_modules", "__pycache__", ".venv", "venv", "site-packages", ".idea", ".vscode"}
# 名字里带这些字样的文件或目录一律不碰（凭据、登录态）
SENSITIVE = re.compile(r"(?i)secret|token|credential|cookie|password|passwd|private[-_ ]?key|^\.env(\..*)?$|\.local\.|"
                       r"profile|login[-_ ]?data|^local state$|\.pem$|\.key$|\.pfx$|\.p12$")
# 浏览器资料目录里的标志文件：目录里有任何一个，整个目录跳过
BROWSER_MARKERS = {"Local State", "Cookies", "Login Data", "Web Data", "Preferences", "Secure Preferences"}

INDEX_FILENAME = "shared-folders.sqlite"
SYNC_INTERVAL = 10 * 60  # 新增和修改的文件，最迟这一轮后台同步后能搜到

_tls = threading.local()
_init_lock = threading.Lock()
_sync_lock = threading.Lock()
_warm_lock = threading.Lock()
_warm_stop = threading.Event()
_warm_thread: threading.Thread | None = None


def folders_for(viewer_key: str | None) -> list[dict]:
    if not viewer_key:
        return []
    return [f for f in config.shared_folders() if viewer_key in f["readers"]]


def _is_link(entry: os.DirEntry | Path) -> bool:
    try:
        if entry.is_symlink():
            return True
        junction = getattr(entry, "is_junction", None)
        return bool(junction and junction())
    except OSError:
        return True


def _excluded(rel: str, patterns: list[str]) -> bool:
    rel_l = rel.lower()
    for p in patterns:
        p = p.replace("\\", "/").strip("/").lower()
        if rel_l == p or rel_l.startswith(p + "/") or fnmatch(rel_l, p):
            return True
    return False


def _name_ok(name: str) -> bool:
    return not SENSITIVE.search(name)


def _file_ok(name: str, size: int) -> bool:
    ext = os.path.splitext(name)[1].lower()
    if not _name_ok(name):
        return False
    if ext in TEXT_TYPES:
        return size <= MAX_TEXT_BYTES
    return ext in DOC_TYPES and size <= MAX_DOC_BYTES


def walk(folder: dict):
    """逐个给出 (相对路径, 绝对路径, stat)；不跟随任何链接。"""
    root = Path(folder["path"])
    if not root.is_dir() or _is_link(root):
        return
    stack = [(root, "")]
    while stack:
        here, prefix = stack.pop()
        try:
            entries = list(os.scandir(here))
        except OSError:
            continue
        if any(e.name in BROWSER_MARKERS for e in entries):
            continue
        for e in entries:
            rel = f"{prefix}{e.name}"
            if _is_link(e) or not _name_ok(e.name) or _excluded(rel, folder["exclude"]):
                continue
            try:
                if e.is_dir(follow_symlinks=False):
                    if e.name not in SKIP_DIRS:
                        stack.append((Path(e.path), rel + "/"))
                elif e.is_file(follow_symlinks=False):
                    st = e.stat(follow_symlinks=False)
                    if _file_ok(e.name, st.st_size):
                        yield rel, Path(e.path), st
            except OSError:
                continue


def resolve(folder: dict, rel: str) -> Path | None:
    """把相对路径安全地落回文件夹里的一个文件；任何一步不合规就返回 None。"""
    if not rel or "\\" in rel or ":" in rel or rel.startswith("/"):
        return None
    parts = rel.split("/")  # 不做规整：只认 walk 给出的标准写法
    if any(p in ("", ".", "..") or p != p.strip() for p in parts):
        return None
    if _excluded(rel, folder["exclude"]):
        return None
    root = Path(folder["path"])
    if _is_link(root):
        return None
    here = root
    for i, part in enumerate(parts):
        if not _name_ok(part) or (i < len(parts) - 1 and part in SKIP_DIRS):
            return None
        try:
            if any((here / m).exists() for m in BROWSER_MARKERS):
                return None
        except OSError:
            return None
        here = here / part
        if _is_link(here):
            return None
    try:
        real, real_root = here.resolve(strict=True), root.resolve(strict=True)
        st = here.stat()
    except OSError:
        return None
    if not real.is_relative_to(real_root) or not here.is_file() or not _file_ok(here.name, st.st_size):
        return None
    return here


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "gb18030"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _html_text(text: str) -> str:
    text = re.sub(r"(?is)<(script|style)\b.*?</\1>", " ", text)
    text = re.sub(r"(?s)<[^>]+>", " ", text)
    return re.sub(r"[ \t\r\f\v]+", " ", html.unescape(text))


def _pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n\n".join((page.extract_text() or "") for page in reader.pages)


def _docx(path: Path) -> str:
    import docx

    d = docx.Document(str(path))
    out = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            out.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(out)


def _pptx(path: Path) -> str:
    from pptx import Presentation

    out = []
    for n, slide in enumerate(Presentation(str(path)).slides, 1):
        out.append(tr("folder.page", n=n))
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                out.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    out.append("\t".join(cell.text for cell in row.cells))
        if slide.has_notes_slide:
            out.append(slide.notes_slide.notes_text_frame.text)
    return "\n".join(out)


def _xlsx(path: Path) -> str:
    import warnings

    from openpyxl import load_workbook

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")  # 缺默认样式之类的提示，和抽文字无关
        wb = load_workbook(str(path), read_only=True, data_only=True)
    out, total = [], 0
    try:
        for ws in wb.worksheets:
            out.append(tr("folder.sheet", name=ws.title))
            for row in ws.iter_rows(values_only=True):
                line = "\t".join("" if v is None else str(v) for v in row).rstrip()
                if line:
                    out.append(line)
                    total += len(line)
                if total > MAX_CHARS:
                    return "\n".join(out)
    finally:
        wb.close()
    return "\n".join(out)


EXTRACT = {".pdf": _pdf, ".docx": _docx, ".pptx": _pptx, ".xlsx": _xlsx}


def text_of(path: Path, st: os.stat_result | None = None) -> str:
    """抽出正文，最多留 MAX_CHARS + 1 个字。抽不出来返回空串。

    不在进程里留正文。同步和 fetch 按索引里的修改时间、大小决定要不要调用这里；
    st 只是调用方已经拿到的 stat，这里不再用它做缓存键。
    """
    del st
    ext = path.suffix.lower()
    try:
        if ext in EXTRACT:
            text = EXTRACT[ext](path)
        else:
            text = _read_text(path)
            if ext in (".html", ".htm"):
                text = _html_text(text)
    except ImportError:
        text = ""
        log.warning("缺少抽取 %s 文字的库，装 onestead[docs] 后才能读", ext)
    except Exception:  # 坏文件、加密 PDF 等：当作读不出文字
        log.info("抽不出文字：%s", path.name, exc_info=True)
        text = ""
    if len(text) > MAX_CHARS + 1:
        text = text[:MAX_CHARS + 1]  # 多留一个字，fetch 据此知道被截断了
    # PDF、JSON 里常有半个 emoji（落单的代理码位），SQLite 存不进去；换成替代字符
    return text.encode("utf-8", "replace").decode("utf-8")


def doc_id(folder: dict, rel: str) -> str:
    return f"folder:{folder['key']}/{rel}"


def parse_id(document_id: str) -> tuple[str, str] | None:
    if not document_id.startswith("folder:"):
        return None
    key, _, rel = document_id.removeprefix("folder:").partition("/")
    return (key, rel) if key and rel else None


def _db_path() -> Path:
    return config.state_dir() / INDEX_FILENAME


def _restrict(path: Path) -> None:
    """索引里是业务文字。POSIX 上数据库和 WAL 只留本人读写；Windows 靠状态目录本身的权限。"""
    if os.name == "nt":
        return
    for suffix in ("", "-wal", "-shm"):
        candidate = Path(str(path) + suffix)
        if candidate.exists():
            os.chmod(candidate, 0o600)


def _prepare_db_file(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if os.name == "nt" or path.exists():
        return
    fd = os.open(path, os.O_CREAT | os.O_RDWR, 0o600)
    os.close(fd)


def _unlink_db(path: Path) -> None:
    for suffix in ("-wal", "-shm", ""):
        candidate = Path(str(path) + suffix)
        try:
            candidate.unlink(missing_ok=True)
        except OSError:
            if suffix == "" and candidate.exists():
                raise


def _close_quiet(conn: sqlite3.Connection) -> None:
    try:
        conn.close()
    except sqlite3.Error:
        pass


def close() -> None:
    """关掉这个线程上的索引连接。下一次同步或搜索会重新打开。"""
    conn = getattr(_tls, "conn", None)
    if conn is not None:
        _close_quiet(conn)
    _tls.conn = None
    _tls.path = None


def _is_corrupt(exc: sqlite3.DatabaseError) -> bool:
    message = str(exc).lower()
    return "malformed" in message or "not a database" in message or "disk image" in message


def _ensure_schema(conn: sqlite3.Connection) -> None:
    conn.execute("""
        CREATE TABLE IF NOT EXISTS files (
            id INTEGER PRIMARY KEY,
            folder_key TEXT NOT NULL,
            rel TEXT NOT NULL,
            mtime_ns INTEGER NOT NULL,
            size INTEGER NOT NULL,
            title TEXT NOT NULL,
            body TEXT NOT NULL,
            UNIQUE (folder_key, rel)
        )
    """)
    conn.execute("""
        CREATE VIRTUAL TABLE IF NOT EXISTS files_fts USING fts5(
            title, body,
            content='files', content_rowid='id',
            tokenize='trigram case_sensitive 0'
        )
    """)
    conn.execute("""
        CREATE TRIGGER IF NOT EXISTS files_ai AFTER INSERT ON files BEGIN
            INSERT INTO files_fts(rowid, title, body) VALUES (new.id, new.title, new.body);
        END
    """)
    conn.execute("""
        CREATE TRIGGER IF NOT EXISTS files_ad AFTER DELETE ON files BEGIN
            INSERT INTO files_fts(files_fts, rowid, title, body)
            VALUES ('delete', old.id, old.title, old.body);
        END
    """)
    conn.execute("""
        CREATE TRIGGER IF NOT EXISTS files_au AFTER UPDATE ON files BEGIN
            INSERT INTO files_fts(files_fts, rowid, title, body)
            VALUES ('delete', old.id, old.title, old.body);
            INSERT INTO files_fts(rowid, title, body) VALUES (new.id, new.title, new.body);
        END
    """)
    files = conn.execute("SELECT count(*) FROM files").fetchone()[0]
    indexed = conn.execute("SELECT count(*) FROM files_fts").fetchone()[0]
    if files != indexed:
        conn.execute("INSERT INTO files_fts(files_fts) VALUES ('rebuild')")


def _open_once(path: Path) -> sqlite3.Connection:
    _prepare_db_file(path)
    conn = sqlite3.connect(path, isolation_level=None)
    try:
        conn.row_factory = sqlite3.Row
        conn.execute("PRAGMA journal_mode=WAL")
        conn.execute("PRAGMA busy_timeout=5000")
        _ensure_schema(conn)
        _restrict(path)
        return conn
    except Exception:
        conn.close()
        raise


def _open(path: Path) -> sqlite3.Connection:
    with _init_lock:
        try:
            return _open_once(path)
        except sqlite3.DatabaseError as exc:
            if not path.exists() or not _is_corrupt(exc):
                raise
            _unlink_db(path)
            return _open_once(path)


def _conn() -> sqlite3.Connection:
    path = _db_path()
    key = str(path)
    conn = getattr(_tls, "conn", None)
    if conn is not None and getattr(_tls, "path", None) == key and path.exists():
        try:
            conn.execute("SELECT 1 FROM files LIMIT 1")
            return conn
        except sqlite3.DatabaseError as exc:
            _close_quiet(conn)
            _tls.conn = None
            if not _is_corrupt(exc):
                raise
    elif conn is not None:
        _close_quiet(conn)
        _tls.conn = None
    conn = _open(path)
    _tls.conn = conn
    _tls.path = key
    return conn


def _write_file(conn: sqlite3.Connection, folder: dict, rel: str, st: os.stat_result, body: str) -> None:
    conn.execute("""
        INSERT INTO files (folder_key, rel, mtime_ns, size, title, body)
        VALUES (?, ?, ?, ?, ?, ?)
        ON CONFLICT(folder_key, rel) DO UPDATE SET
            mtime_ns=excluded.mtime_ns,
            size=excluded.size,
            title=excluded.title,
            body=excluded.body
    """, (folder["key"], rel, st.st_mtime_ns, st.st_size, f"{folder['name']}/{rel}", body))


def _sync_folder(conn: sqlite3.Connection, folder: dict) -> tuple[int, int]:
    key = folder["key"]
    existing = {
        row["rel"]: (row["mtime_ns"], row["size"], row["title"])
        for row in conn.execute(
            "SELECT rel, mtime_ns, size, title FROM files WHERE folder_key=?", (key,)
        )
    }
    seen: set[str] = set()
    extracted = 0
    for rel, path, st in walk(folder):
        seen.add(rel)
        title = f"{folder['name']}/{rel}"
        prev = existing.get(rel)
        if prev and prev[0] == st.st_mtime_ns and prev[1] == st.st_size:
            if prev[2] != title:
                conn.execute("UPDATE files SET title=? WHERE folder_key=? AND rel=?", (title, key, rel))
            continue
        try:
            _write_file(conn, folder, rel, st, text_of(path, st))
        except sqlite3.DatabaseError as exc:
            if _is_corrupt(exc):
                raise
            log.warning("索引写不进这个文件，先跳过：%s（%s）", rel, exc)  # 一个怪文件不拖住整轮同步
            continue
        extracted += 1
    stale = [rel for rel in existing if rel not in seen]
    for rel in stale:
        conn.execute("DELETE FROM files WHERE folder_key=? AND rel=?", (key, rel))
    return extracted, len(stale)


def _drop_removed_folders(conn: sqlite3.Connection, keys: list[str]) -> int:
    if not keys:
        count = conn.execute("SELECT count(*) FROM files").fetchone()[0]
        conn.execute("DELETE FROM files")
        return count
    marks = ",".join("?" * len(keys))
    count = conn.execute(
        f"SELECT count(*) FROM files WHERE folder_key NOT IN ({marks})", keys
    ).fetchone()[0]
    conn.execute(f"DELETE FROM files WHERE folder_key NOT IN ({marks})", keys)
    return count


def _sync_body(viewer_keys: set[str] | None) -> None:
    configured = config.shared_folders()
    targets = configured
    if viewer_keys is not None:
        targets = [folder for folder in configured if viewer_keys & set(folder["readers"])]
    conn = _conn()
    extracted = deleted = 0
    for folder in targets:
        got, removed = _sync_folder(conn, folder)
        extracted += got
        deleted += removed
    deleted += _drop_removed_folders(conn, [folder["key"] for folder in configured])
    _restrict(_db_path())
    if extracted or deleted:
        log.info("共享文件夹索引已同步，更新 %s 个，删除 %s 个", extracted, deleted)


def sync(viewer_keys: set[str] | None = None) -> None:
    """按 walk 的过滤规则把索引补到和文件夹一致。同一时刻只有一个同步。"""
    with _sync_lock:
        _sync_body(viewer_keys)


def _sync_if_idle() -> None:
    if not _sync_lock.acquire(blocking=False):
        return
    try:
        _sync_body(None)
    finally:
        _sync_lock.release()


def shutdown() -> None:
    """停掉后台同步，并关掉这个线程的索引连接。"""
    global _warm_thread
    _warm_stop.set()
    thread = _warm_thread
    if thread is not None and thread.is_alive() and thread is not threading.current_thread():
        thread.join(timeout=5)
    with _warm_lock:
        if _warm_thread is thread and (thread is None or not thread.is_alive()):
            _warm_thread = None
    close()


def warm(viewer_keys: set[str] | None = None) -> None:
    """启动时后台同步一次，之后每 10 分钟再同步。正在同步时搜索不进这个锁。"""
    global _warm_thread

    def loop() -> None:
        try:
            while not _warm_stop.is_set():
                try:
                    sync(viewer_keys)
                except Exception:
                    log.exception("共享文件夹索引同步失败")
                if _warm_stop.wait(SYNC_INTERVAL):
                    break
        finally:
            close()

    with _warm_lock:
        if _warm_thread and _warm_thread.is_alive():
            return
        _warm_stop.clear()
        _warm_thread = threading.Thread(target=loop, daemon=True, name="shared-folders-index")
        _warm_thread.start()


def _like_pattern(term: str) -> str:
    escaped = term.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
    return f"%{escaped}%"


def _fts_expr(terms: list[str]) -> str | None:
    parts = []
    for term in terms:
        if len(term) < 3:
            continue
        quoted = '"' + term.replace('"', '""') + '"'
        parts.append(f"(title : {quoted} OR body : {quoted})")
    return " AND ".join(parts) if parts else None


def _select(conn: sqlite3.Connection, keys: list[str], terms: list[str], limit: int, *, use_fts: bool):
    where = [f"folder_key IN ({','.join('?' * len(keys))})"]
    params: list = list(keys)
    fts_expr = _fts_expr(terms) if use_fts else None
    if fts_expr:
        where.append("id IN (SELECT rowid FROM files_fts WHERE files_fts MATCH ?)")
        params.append(fts_expr)
    for term in terms:
        # 短词 trigram 对不上；长词也用 LIKE 再确认一遍，避免分词误命中。% 和 _ 是字面量。
        where.append("(lower(title) LIKE ? ESCAPE '\\' OR lower(body) LIKE ? ESCAPE '\\')")
        pattern = _like_pattern(term)
        params.extend((pattern, pattern))
    if terms:
        rank = "CASE WHEN " + " AND ".join("lower(title) LIKE ? ESCAPE '\\'" for _ in terms) + " THEN 0 ELSE 1 END"
        params.extend(_like_pattern(term) for term in terms)
        order = rank + ", title"
    else:
        order = "title"
    params.append(limit)
    sql = "SELECT folder_key, rel FROM files WHERE " + " AND ".join(where) + f" ORDER BY {order} LIMIT ?"
    return conn.execute(sql, params).fetchall()


def _query(conn: sqlite3.Connection, keys: list[str], terms: list[str], limit: int):
    if any(len(term) >= 3 for term in terms):
        try:
            return _select(conn, keys, terms, limit, use_fts=True)
        except sqlite3.OperationalError as exc:
            if _is_corrupt(exc):
                raise
            log.info("全文检索语句无法执行，改为逐字匹配")
            return _select(conn, keys, terms, limit, use_fts=False)
    return _select(conn, keys, terms, limit, use_fts=False)


def _search(viewer_key: str | None, terms: list[str], limit: int) -> list[dict]:
    visible = folders_for(viewer_key)
    if not visible:
        return []
    rows = _query(_conn(), [folder["key"] for folder in visible], [term.lower() for term in terms], limit)
    by_key = {folder["key"]: folder for folder in visible}
    return [_row(by_key[row["folder_key"]], row["rel"], "") for row in rows]


def _rebuild_and_search(viewer_key: str | None, terms: list[str], limit: int) -> list[dict]:
    close()
    _unlink_db(_db_path())
    _sync_body(None)
    return _search(viewer_key, terms, limit)


def search(viewer_key: str | None, terms: list[str], limit: int = 50) -> list[dict]:
    """所有词都要出现在路径或正文里，不分大小写；路径命中的排前面。"""
    thread = _warm_thread
    if thread is None or not thread.is_alive():
        _sync_if_idle()  # 没有后台同步（命令行、测试）才在搜索前补一次；服务里交给后台，搜索不等遍历
    try:
        return _search(viewer_key, terms, limit)
    except sqlite3.DatabaseError as exc:
        if not _is_corrupt(exc):
            raise
        if not _sync_lock.acquire(blocking=False):
            return []
        try:
            return _rebuild_and_search(viewer_key, terms, limit)
        finally:
            _sync_lock.release()


def _indexed_body(folder: dict, rel: str, path: Path, st: os.stat_result) -> str:
    conn = _conn()
    row = conn.execute(
        "SELECT mtime_ns, size, body FROM files WHERE folder_key=? AND rel=?",
        (folder["key"], rel),
    ).fetchone()
    if row is not None and row["mtime_ns"] == st.st_mtime_ns and row["size"] == st.st_size:
        return row["body"]
    body = text_of(path, st)
    _write_file(conn, folder, rel, st, body)
    _restrict(_db_path())
    return body


def fetch(viewer_key: str | None, document_id: str) -> dict | None:
    parsed = parse_id(document_id)
    if not parsed:
        return None
    key, rel = parsed
    folder = next((f for f in folders_for(viewer_key) if f["key"] == key), None)
    path = resolve(folder, rel) if folder else None
    if path is None:
        return None
    text = _indexed_body(folder, rel, path, path.stat())
    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS] + "\n\n" + tr("folder.truncated", n=MAX_CHARS)
    if not text.strip():
        text = tr("folder.no_text")
    return _row(folder, rel, text)


def _row(folder: dict, rel: str, text: str) -> dict:
    return {"id": doc_id(folder, rel), "title": f"{folder['name']}/{rel}", "text": text,
            "metadata": {"type": tr("folder.type"), "folder": folder["name"], "source": "Onestead"}}
